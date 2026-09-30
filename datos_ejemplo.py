"""Llena la base de datos con talentos y evaluaciones de ejemplo, o los borra.

Uso (en la carpeta del proyecto):
    .venv\\Scripts\\python.exe datos_ejemplo.py vista [--cantidad 60] [--semilla N]
    .venv\\Scripts\\python.exe datos_ejemplo.py cargar [--cantidad 60] [--semilla N]
    .venv\\Scripts\\python.exe datos_ejemplo.py borrar

"vista" guarda en un Excel los datos que se cargarían, sin escribir en la BD; con
la misma cantidad y semilla, "cargar" inserta exactamente esos datos.

Solo es para desarrollo y demostraciones (ver los paneles con datos). Cada talento
de ejemplo deja su "alta" en la bitácora con origen "datos_ejemplo.py"; así se
reconocen al borrarlos (no por el correo, que puede coincidir con talentos reales).
Sus evaluaciones y eventos se eliminan en cascada; la bitácora se conserva. Las
evaluaciones se generan sin la IA, con valores coherentes.
"""

import argparse
import datetime
import random
import sys
import unicodedata

from backend import bitacora, catalogos, db, errores, evaluaciones, ia, permisos, talentos

DOMINIO = "@correo.com"
ORIGEN = "datos_ejemplo.py"
SEMILLA = 2026
ARCHIVO_VISTA = "datos_ejemplo_vista.xlsx"
NIVELES = ("bajo", "medio", "alto")
CRITERIOS = ("resultados_cap_ev", "experiencia_lab_ev", "objetivos_fut_ev", "competencias_ev", "liderazgo_prev_ev")

NOMBRES = [
    "José", "Juan", "Luis", "Carlos", "Miguel", "Jorge", "Alejandro", "Fernando", "Ricardo", "Eduardo",
    "Francisco", "Roberto", "Daniel", "Sergio", "Javier", "Manuel", "Antonio", "Raúl", "Óscar", "Diego",
    "Juan Carlos", "José Luis", "Luis Ángel", "Emiliano", "Santiago",
    "María", "Guadalupe", "Juana", "Verónica", "Leticia", "Patricia", "Gabriela", "Alejandra", "Adriana",
    "Claudia", "Laura", "Mónica", "Andrea", "Daniela", "Fernanda", "Rosa", "Silvia", "Ximena", "Paola",
    "Karla", "María Fernanda", "María José", "Ana Laura", "Diana", "Valeria",
]
APELLIDOS = [
    "Hernández", "García", "Martínez", "López", "González", "Pérez", "Rodríguez", "Sánchez", "Ramírez",
    "Cruz", "Flores", "Gómez", "Morales", "Vázquez", "Jiménez", "Reyes", "Díaz", "Torres", "Gutiérrez",
    "Ruiz", "Mendoza", "Aguilar", "Ortiz", "Moreno", "Castillo", "Romero", "Álvarez", "Méndez", "Chávez",
    "Rivera", "Juárez", "Ramos", "Domínguez", "Herrera", "Medina", "Castro", "Vargas", "Guzmán", "Velázquez",
    "Salazar", "Rojas", "Contreras", "Luna", "Ortega", "Soto", "Navarro", "Solano", "Estrada", "Delgado",
]

def _telefonos():
    # Orden visual: 1112345678, 1111234567, ... 1111111111; luego 2111111111,
    # 2211111111, ... 2222222222; 3222222222, ... hasta 9999999999 (88 números)
    lista = ["1" * unos + "23456789"[:10 - unos] for unos in range(3, 11)]
    for digito in range(2, 10):
        lista += [str(digito) * cantidad + str(digito - 1) * (10 - cantidad) for cantidad in range(1, 11)]
    return lista


TELEFONOS = _telefonos()

COMPETENCIAS = {
    "Desarrollo": ["Python, SQL y control de versiones con Git", "Java, Spring Boot y pruebas unitarias",
                   "JavaScript, React y diseño de APIs REST", "Arquitectura de microservicios y revisión de código"],
    "Calidad": ["Diseño de casos de prueba y reporte de defectos", "Automatización de pruebas con Selenium",
                "Pruebas de regresión y de rendimiento", "Normas ISO 9001 y mejora de procesos"],
    "Soporte técnico": ["Atención a usuarios y resolución de incidentes", "Administración de Windows y Office 365",
                        "Diagnóstico de hardware y redes locales", "Documentación de soluciones en la mesa de ayuda"],
    "Infraestructura/IT/DevOps": ["Linux, Docker y Kubernetes", "Integración continua con GitHub Actions",
                                  "Administración de servidores y respaldos", "Monitoreo, redes y seguridad perimetral"],
    "Ventas/marketing": ["Negociación y cierre de ventas B2B", "Marketing digital y redes sociales",
                         "Manejo de CRM y seguimiento de clientes", "Análisis de mercado y presentaciones comerciales"],
    "Finanzas": ["Contabilidad general y conciliaciones bancarias", "Excel avanzado y modelos financieros",
                 "Presupuestos, costos y flujo de efectivo", "Cumplimiento fiscal y facturación electrónica"],
    "Recursos Humanos": ["Reclutamiento y selección de personal", "Nómina, IMSS y prestaciones",
                         "Clima laboral y comunicación interna", "Capacitación y evaluación del desempeño"],
}
CAPACITACION = {
    "bajo": ["Curso de inducción de la empresa.", "Taller básico de ofimática; no ha concluido otros cursos.",
             "Curso en línea de atención al cliente, sin certificación."],
    "medio": ["Diplomado en su área concluido en 2023 y dos cursos en línea con constancia.",
              "Curso de comunicación efectiva y certificación básica de su especialidad.",
              "Talleres internos de trabajo en equipo y un curso técnico de 40 horas."],
    "alto": ["Certificación profesional vigente, diplomado en liderazgo y cursos técnicos avanzados.",
             "Maestría en curso, dos certificaciones internacionales y mentoría de nuevos ingresos.",
             "Especialización concluida con honores y participación como instructor interno."],
}
EXPERIENCIA = {
    "bajo": ["Primer empleo formal; prácticas profesionales de seis meses.",
             "Un año de experiencia en un puesto operativo similar."],
    "medio": ["Tres años en puestos similares en empresas medianas.",
              "Cuatro años de experiencia, dos de ellos en el sector tecnológico."],
    "alto": ["Más de ocho años de experiencia, con proyectos a nivel nacional.",
             "Seis años en empresas del sector, con participación en proyectos internacionales."],
}
OBJETIVOS = {
    "bajo": ["Mantenerse en su puesto actual.", "Aún no define sus metas profesionales."],
    "medio": ["Obtener una certificación en su área en el próximo año.", "Especializarse en su área actual."],
    "alto": ["Asumir la coordinación de un equipo en los próximos dos años.",
             "Llegar a un puesto de gerencia y formar a nuevos talentos."],
}
LIDERAZGO = {
    "bajo": ["Sin experiencia de liderazgo.", "Participa en equipos sin tareas de coordinación."],
    "medio": ["Ha coordinado proyectos pequeños de su área.", "Apoya en la integración de personal nuevo."],
    "alto": ["Lideró un equipo de cinco personas durante un año.",
             "Coordina proyectos entre áreas y es referente técnico del equipo."],
}
FORTALEZAS = {
    "bajo": "cumple con las tareas asignadas y muestra disposición para aprender",
    "medio": "tiene una base sólida en su área y resultados constantes",
    "alto": "destaca por sus resultados, su preparación y su capacidad de influir en el equipo",
}
RECOMENDACIONES = {
    "bajo": "Se recomienda un plan de desarrollo con metas mensuales, acompañamiento cercano y capacitación básica.",
    "medio": "Se recomienda asignarle proyectos con mayor responsabilidad y una certificación de su especialidad.",
    "alto": "Se recomienda incluirlo en el plan de sucesión y darle participación en la toma de decisiones.",
}


def sin_acentos(texto):
    normal = unicodedata.normalize("NFKD", texto)
    return "".join(letra for letra in normal if not unicodedata.combining(letra))


def _nivel_con_ruido(nivel, rng):
    # Devuelve el mismo nivel la mayoría de las veces, o uno vecino
    indice = NIVELES.index(nivel) + rng.choices((-1, 0, 1), weights=(2, 6, 2))[0]
    return NIVELES[min(max(indice, 0), 2)]


def _desempeno(rng, indice):
    # Tercios: muy bajos (5-30), medios (31-70) y altos (71-100)
    minimo, maximo = ((5, 30), (31, 70), (71, 100))[indice % 3]
    return rng.randint(minimo, maximo)


def _fecha_aleatoria(rng, desde, hasta):
    return desde + datetime.timedelta(days=rng.randint(0, (hasta - desde).days))


def _talento(rng, nombre, correo, area, fecha_ingreso, desempeno):
    nivel = ia.nivel_de_puntuacion(desempeno)
    datos = {
        "nombre": nombre,
        "area": area,
        "correo": correo,
        "tel_mx": None,  # se asigna en orden al final de generar()
        "fecha_ingreso": fecha_ingreso,
        "desempeno_previo": desempeno,
        "competencias": rng.choice(COMPETENCIAS[area]),
        "capacitacion": rng.choice(CAPACITACION[_nivel_con_ruido(nivel, rng)]),
        "experiencia": rng.choice(EXPERIENCIA[_nivel_con_ruido(nivel, rng)]),
        "objetivos_futuros": rng.choice(OBJETIVOS[_nivel_con_ruido(nivel, rng)]),
        "liderazgo_previo": rng.choice(LIDERAZGO[_nivel_con_ruido(nivel, rng)]),
    }
    # Algunos indicadores quedan sin capturar, como pasaría en la realidad
    for campo in talentos.INDICADORES:
        if campo != "competencias" and rng.random() < 0.08:
            datos[campo] = talentos.SIN_REGISTRO
    return datos


def _evaluacion(rng, talento, puntuacion, fecha):
    potencial = ia.nivel_de_puntuacion(puntuacion)
    evaluacion = {
        "fecha_evaluacion": fecha,
        "puntuacion_talento_por_ia": puntuacion,
        "potencial_crecimiento_nivel": potencial,
        "desempeno_pre_ev": ia.nivel_de_puntuacion(talento["desempeno_previo"]),
        **{criterio: _nivel_con_ruido(potencial, rng) for criterio in CRITERIOS},
    }
    evaluacion["clasificacion_9box"] = evaluaciones.clasificar_9box(talento["desempeno_previo"], potencial)
    evaluacion["reporte_ia"] = (
        f"{talento['nombre']} presenta un potencial de crecimiento {potencial} ({puntuacion} de 100 puntos), "
        f"con un desempeño previo de {talento['desempeno_previo']} de 100. En general {FORTALEZAS[potencial]}. "
        f"Su clasificación en la matriz 9-box es \"{evaluacion['clasificacion_9box']}\". "
        f"{RECOMENDACIONES[potencial]} (Reporte de ejemplo, no generado por la IA.)"
    )
    return evaluacion


def _evaluaciones(rng, talento, hoy):
    # De 0 a 4 evaluaciones separadas de 1 a 4 meses, dentro de los últimos 18 meses
    cantidad = rng.choices((0, 1, 2, 3, 4), weights=(1, 3, 3, 2, 1))[0]
    desde = max(talento["fecha_ingreso"] + datetime.timedelta(days=30), hoy - datetime.timedelta(days=540))
    if cantidad == 0 or desde > hoy:
        return []
    # El potencial se inclina hacia el nivel del desempeño, pero cubre las 9 celdas de la 9-box
    pesos = {"bajo": (4, 4, 2), "medio": (3, 4, 3), "alto": (3, 3, 4)}[ia.nivel_de_puntuacion(talento["desempeno_previo"])]
    potencial = rng.choices(NIVELES, weights=pesos)[0]
    puntuacion = rng.randint(*{"bajo": (12, 30), "medio": (31, 70), "alto": (71, 96)}[potencial])
    fecha = _fecha_aleatoria(rng, desde, desde + (hoy - desde) // 2)
    lista = []
    for _ in range(cantidad):
        if fecha > hoy:
            break
        lista.append(_evaluacion(rng, talento, puntuacion, fecha))
        puntuacion = min(max(puntuacion + rng.randint(-8, 12), 8), 98)
        fecha += datetime.timedelta(days=rng.randint(30, 120))
    return lista


def _eventos(rng, talento, lista, hoy):
    eventos = [("ingreso", f"Ingreso al área {talento['area']}.", talento["fecha_ingreso"])]
    if rng.random() < 0.3:
        eventos.append(("capacitacion", "Concluyó un curso de capacitación de su área.",
                        _fecha_aleatoria(rng, talento["fecha_ingreso"], hoy)))
    anterior = None
    for evaluacion in lista:
        nivel = evaluacion["potencial_crecimiento_nivel"]
        eventos.append(("evaluacion_ia",
                        f"Evaluación con IA: potencial {nivel}, {evaluacion['puntuacion_talento_por_ia']} puntos, "
                        f"{evaluacion['clasificacion_9box']}.", evaluacion["fecha_evaluacion"]))
        if anterior is not None and anterior != nivel:
            eventos.append(("cambio_clasificacion", f"Potencial de crecimiento: {anterior} -> {nivel}.",
                            evaluacion["fecha_evaluacion"]))
        anterior = nivel
    return eventos


def generar(cantidad, rng, hoy=None, ocupados=()):
    if not 1 <= cantidad <= len(TELEFONOS):
        raise ValueError(f"La cantidad debe estar entre 1 y {len(TELEFONOS)}.")
    # Lista de talentos de ejemplo, cada uno con sus evaluaciones y eventos;
    # "ocupados" son los correos que ya existen en la BD
    hoy = hoy or datetime.date.today()
    usados_nombres, usados_correos, generados = set(), set(ocupados), []
    while len(generados) < cantidad:
        primero, segundo = rng.sample(APELLIDOS, 2)
        nombre = f"{rng.choice(NOMBRES)} {primero} {segundo}"
        if nombre in usados_nombres:
            continue
        usados_nombres.add(nombre)
        base = sin_acentos(f"{nombre.split()[0]}.{primero}").lower()
        correo, numero = base + DOMINIO, 1
        while correo in usados_correos:
            numero += 1
            correo = f"{base}{numero}{DOMINIO}"
        usados_correos.add(correo)

        area = catalogos.AREAS[len(generados) % len(catalogos.AREAS)]
        fecha_ingreso = _fecha_aleatoria(rng, datetime.date(2020, 1, 1), hoy)
        talento = _talento(rng, nombre, correo, area, fecha_ingreso, _desempeno(rng, len(generados)))
        talento["evaluaciones"] = _evaluaciones(rng, talento, hoy)
        talento["eventos"] = _eventos(rng, talento, talento["evaluaciones"], hoy)
        generados.append(talento)
    rng.shuffle(generados)
    for talento, telefono in zip(generados, TELEFONOS):
        talento["tel_mx"] = telefono
    return generados


def ids_de_ejemplo(cursor):
    cursor.execute(
        "SELECT t.id_talento FROM Talentos t JOIN Bitacora b ON b.tabla_afectada = 'Talentos' "
        "AND b.id_registro_afectado = t.id_talento AND b.accion = 'alta' "
        "WHERE JSON_UNQUOTE(JSON_EXTRACT(b.detalle, '$.origen')) = %s",
        (ORIGEN,),
    )
    return [fila["id_talento"] for fila in cursor.fetchall()]


def guardar(cursor, lista, id_usuario):
    columnas_eval = ("fecha_evaluacion", *evaluaciones.COLUMNAS)
    for talento in lista:
        cursor.execute(
            "INSERT INTO Talentos (" + ", ".join(talentos.CAMPOS) + ", fecha_modificacion) "
            "VALUES (" + ", ".join(["%s"] * len(talentos.CAMPOS)) + ", %s)",
            (*(talento[campo] for campo in talentos.CAMPOS), talento["fecha_ingreso"]),
        )
        id_talento = cursor.lastrowid
        bitacora.registrar(cursor, id_usuario, "alta", "Talentos", id_talento, {"origen": ORIGEN})
        for evaluacion in talento["evaluaciones"]:
            cursor.execute(
                "INSERT INTO Evaluaciones (id_talento, " + ", ".join(columnas_eval) + ") "
                "VALUES (%s, " + ", ".join(["%s"] * len(columnas_eval)) + ")",
                (id_talento, *(evaluacion[columna] for columna in columnas_eval)),
            )
        cursor.executemany(
            "INSERT INTO Eventos_talento (id_talento, tipo_evento, descripcion, fecha_evento, "
            "id_usuario_registro) VALUES (%s, %s, %s, %s, %s)",
            [(id_talento, tipo, descripcion, fecha, id_usuario) for tipo, descripcion, fecha in talento["eventos"]],
        )


def _correos_ocupados(cursor):
    cursor.execute("SELECT correo FROM Talentos")
    return {fila["correo"] for fila in cursor.fetchall()}


def vista(cantidad, semilla=SEMILLA, archivo=ARCHIVO_VISTA):
    # Guarda en Excel los datos que se cargarían (solo lee los correos de la BD)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    with db.transaccion() as cursor:
        lista = generar(cantidad, random.Random(semilla), ocupados=_correos_ocupados(cursor))
    libro = Workbook()
    hojas = {
        "Talentos": ([*talentos.CAMPOS, "evaluaciones"],
                     [[t[c] for c in talentos.CAMPOS] + [len(t["evaluaciones"])] for t in lista]),
        "Evaluaciones": (["nombre", "fecha_evaluacion", *evaluaciones.COLUMNAS],
                         [[t["nombre"], e["fecha_evaluacion"], *(e[c] for c in evaluaciones.COLUMNAS)]
                          for t in lista for e in t["evaluaciones"]]),
        "Eventos": (["nombre", "tipo_evento", "descripcion", "fecha_evento"],
                    [[t["nombre"], *evento] for t in lista for evento in t["eventos"]]),
    }
    for indice, (titulo, (encabezados, filas)) in enumerate(hojas.items()):
        hoja = libro.active if indice == 0 else libro.create_sheet()
        hoja.title = titulo
        hoja.append(encabezados)
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="0B1F3A")
        for fila in filas:
            hoja.append(fila)
            for celda in hoja[hoja.max_row]:
                if isinstance(celda.value, datetime.datetime):
                    celda.number_format = "yyyy-mm-dd"
        for columna in hoja.columns:
            ancho = max(len(str(celda.value or "")[:10] if isinstance(celda.value, datetime.date)
                            else str(celda.value or "")) for celda in columna)
            hoja.column_dimensions[columna[0].column_letter].width = min(ancho + 2, 60)
        hoja.freeze_panes = "A2"
    libro.save(archivo)
    return len(lista), sum(len(t["evaluaciones"]) for t in lista)


def cargar(cantidad, semilla=SEMILLA):
    # Devuelve (talentos cargados, evaluaciones cargadas) o lanza ValueError
    with db.transaccion() as cursor:
        if ids_de_ejemplo(cursor):
            raise ValueError("Ya hay talentos de ejemplo en la base de datos. Bórrelos primero con 'borrar'.")
        cursor.execute("SELECT id_usuario FROM Usuarios WHERE cargo = %s AND activo ORDER BY id_usuario LIMIT 1",
                       (permisos.ADMINISTRADOR,))
        administrador = cursor.fetchone()
        if administrador is None:
            raise ValueError("No hay ningún Administrador de RH activo en la base de datos.")
        lista = generar(cantidad, random.Random(semilla), ocupados=_correos_ocupados(cursor))
        guardar(cursor, lista, administrador["id_usuario"])
    return len(lista), sum(len(talento["evaluaciones"]) for talento in lista)


def borrar():
    # Devuelve cuántos talentos de ejemplo se borraron (con sus evaluaciones y eventos)
    with db.transaccion() as cursor:
        ids = ids_de_ejemplo(cursor)
        if ids:
            cursor.execute("DELETE FROM Talentos WHERE id_talento IN (" + ", ".join(["%s"] * len(ids)) + ")",
                           tuple(ids))
        return len(ids)


def main():
    argumentos = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    acciones = argumentos.add_subparsers(dest="accion", required=True)
    for accion, ayuda in (("vista", f"guarda en {ARCHIVO_VISTA} los datos que se cargarían, sin tocar la BD"),
                          ("cargar", "agrega talentos y evaluaciones de ejemplo")):
        opciones_carga = acciones.add_parser(accion, help=ayuda)
        opciones_carga.add_argument("--cantidad", type=int, default=60,
                                    help=f"número de talentos (1-{len(TELEFONOS)}, 60 por omisión)")
        opciones_carga.add_argument("--semilla", type=int, default=SEMILLA,
                                    help=f"otra semilla genera otros datos ({SEMILLA} por omisión)")
    acciones.add_parser("borrar", help="borra solo los talentos cargados con este script")
    opciones = argumentos.parse_args()

    try:
        if opciones.accion in ("vista", "cargar") and not 1 <= opciones.cantidad <= len(TELEFONOS):
            print(f"La cantidad debe estar entre 1 y {len(TELEFONOS)}.")
            return 1
        if opciones.accion == "vista":
            total, evaluadas = vista(opciones.cantidad, opciones.semilla)
            print(f"Se guardaron {total} talentos y {evaluadas} evaluaciones en {ARCHIVO_VISTA} "
                  "(no se escribió nada en la base de datos).")
            return 0
        if opciones.accion == "cargar":
            total, evaluadas = cargar(opciones.cantidad, opciones.semilla)
            print(f"Se cargaron {total} talentos de ejemplo con {evaluadas} evaluaciones.")
            return 0

        with db.transaccion() as cursor:
            total = len(ids_de_ejemplo(cursor))
        if not total:
            print("No hay talentos de ejemplo en la base de datos.")
            return 0
        print(f"Se borrarán {total} talentos de ejemplo, con sus evaluaciones y eventos.")
        if input("¿Continuar? (s/n): ").strip().lower() != "s":
            print("Cancelado; no se hizo ningún cambio.")
            return 1
        print(f"Se borraron {borrar()} talentos de ejemplo.")
        return 0
    except ValueError as problema:
        print(problema)
        return 1
    except errores.ErrorSistema as error:
        print(f"{error.mensaje} (código {error.codigo})")
        return 1


if __name__ == "__main__":
    sys.exit(main())
