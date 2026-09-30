import datetime
import io
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from backend import analitica, db, errores, evaluaciones, ia, permisos, reportes, talentos

GENERADO = datetime.datetime(2026, 9, 30, 10, 15)
USUARIO = {"id_usuario": 1, "nombre": "Prueba RH", "cargo": permisos.ADMINISTRADOR, "area": "Recursos Humanos"}
TALENTO = {"nombre": "Ana García", "area": "Infraestructura/IT/DevOps", "correo": "ana.reporte@ejemplo.com",
           "tel_mx": "5587654321", "fecha_ingreso": datetime.date(2023, 5, 2), "desempeno_previo": 80,
           "competencias": "Comunicación", "capacitacion": "Scrum", "experiencia": "5 años",
           "objetivos_futuros": "Liderar un equipo", "liderazgo_previo": ""}
EVALUACION = {"desempeno_pre_ev": "alto", "resultados_cap_ev": "medio", "experiencia_lab_ev": "alto",
              "objetivos_fut_ev": "medio", "competencias_ev": "alto", "liderazgo_prev_ev": "bajo",
              "potencial_crecimiento_nivel": "alto", "puntuacion_talento_por_ia": 85,
              "reporte_ia": "Ana García muestra un desempeño “sobresaliente” – con liderazgo…",
              "clasificacion_9box": "Alta estrella"}


def _individual(con_evaluacion=True):
    talento = {"id_talento": 7, **TALENTO, "fecha_modificacion": datetime.date(2026, 9, 1)}
    lista = [{"id_evaluacion": 1, "fecha_evaluacion": datetime.date(2026, 3, 1), **EVALUACION,
              "puntuacion_talento_por_ia": 60, "potencial_crecimiento_nivel": "medio",
              "clasificacion_9box": "Profesional clave"},
             {"id_evaluacion": 2, "fecha_evaluacion": datetime.date(2026, 9, 1), **EVALUACION}] if con_evaluacion else []
    eventos = [{"id_evento": 1, "tipo_evento": "ingreso", "descripcion": "Ingreso al área.",
                "fecha_evento": datetime.date(2023, 5, 2)}]
    return {"tipo": "individual", "talento": talento, "evaluaciones": lista, "eventos": eventos,
            "ultima": lista[-1] if lista else None}


def _fila(id_talento, area, nivel=None, puntuacion=None, fecha=None):
    return {"id_talento": id_talento, "nombre": f"Talento {id_talento}", "area": area,
            "fecha_ingreso": datetime.date(2022, 1, 1), "desempeno_previo": 80,
            "id_evaluacion": id_talento if nivel else None, "fecha_evaluacion": fecha,
            "puntuacion_talento_por_ia": puntuacion, "potencial_crecimiento_nivel": nivel,
            "clasificacion_9box": evaluaciones.clasificar_9box(80, nivel) if nivel else None,
            **{llave: nivel for llave in ia.CRITERIOS}}


@pytest.fixture
def general():
    datos = pd.DataFrame([_fila(i, "Calidad" if i % 2 else "Finanzas", ("bajo", "medio", "alto")[i % 3],
                                (20, 50, 80)[i % 3] + i, datetime.date(2026, 9 - i % 3, 10)) for i in range(1, 13)]
                         + [_fila(13, "Finanzas")])
    todas = datos.dropna(subset=["id_evaluacion"])[["id_talento", "fecha_evaluacion", "puntuacion_talento_por_ia"]]
    resumen = reportes.resumen_general(datos, todas, None, datetime.date(2026, 9, 30))
    talentos_ = [{**{c: fila[c] for c in ("id_talento", "nombre", "area", "fecha_ingreso", "desempeno_previo")},
                  "correo": "x@ejemplo.com", "tel_mx": "5512345678", "competencias": "-", "capacitacion": "-",
                  "experiencia": "-", "objetivos_futuros": "-", "liderazgo_previo": "-",
                  "fecha_modificacion": datetime.date(2026, 1, 1), **reportes._ultima_de(fila)}
                 for _, fila in datos.iterrows()]
    filtros = {"area": "Infraestructura/IT/DevOps", "niveles": ["alto"], "clasificaciones": [], "fechas": None}
    return {"tipo": "general", "filtros": filtros, **resumen, "talentos": talentos_}


def test_latin1():
    assert reportes.latin1("“Hola” – ‘sí’ … → ñ") == '"Hola" - \'sí\' ... -> ñ'
    assert reportes.latin1("日本") == "??"


def test_nombres_de_archivo(general):
    assert reportes.nombre_archivo(_individual(), ".pdf", GENERADO) == "Reporte_7_Ana_Garcia_20260930.pdf"
    assert reportes.nombre_archivo(general, ".xlsx", GENERADO) == \
        "Reporte_general_Infraestructura_IT_DevOps_20260930.xlsx"
    sin_area = {**general, "filtros": {**general["filtros"], "area": None}}
    assert reportes.nombre_archivo(sin_area, ".pdf", GENERADO) == "Reporte_general_todas_20260930.pdf"


def test_resumen_general(general):
    indicadores = general["indicadores"]
    assert indicadores["total"] == 13 and indicadores["sin_evaluar"] == 1
    assert indicadores["porcentajes"] == {"bajo": 33.3, "medio": 33.3, "alto": 33.3}
    # Sin filtro de fechas, el periodo es el mes en curso: 4 evaluaciones de septiembre
    assert indicadores["periodo"] == (datetime.date(2026, 9, 1), datetime.date(2026, 9, 30))
    assert indicadores["evaluaciones_periodo"] == 4
    assert sum(d["total"] for d in general["distribucion"]) == 12
    assert [t["puntuacion"] for t in general["top"]] == sorted((t["puntuacion"] for t in general["top"]), reverse=True)
    assert len(general["top"]) == 10
    assert sum(general["conteo_9box"].values()) == 12


def test_pdf_individual_y_general(general):
    for datos in (_individual(), _individual(con_evaluacion=False)):
        contenido = reportes.pdf_individual(datos, USUARIO, GENERADO)
        assert contenido.startswith(b"%PDF")
    assert reportes.pdf_general(general, USUARIO, GENERADO).startswith(b"%PDF")
    vacio = {**general, "distribucion": [], "top": []}
    assert reportes.pdf_general(vacio, USUARIO, GENERADO).startswith(b"%PDF")


def _libro(contenido):
    return load_workbook(io.BytesIO(contenido))


def test_excel_individual():
    libro = _libro(reportes.excel_individual(_individual(), USUARIO, GENERADO))
    assert libro.sheetnames == ["Talento", "Evaluaciones", "Reporte IA"]
    for hoja in libro.worksheets:
        assert "30/09/2026 10:15" in hoja["A1"].value
        assert hoja["A2"].font.bold and hoja["A2"].fill.fgColor.rgb.endswith("D9D9D9")
        assert hoja.freeze_panes == "A3" and hoja.auto_filter.ref.startswith("A2:")
    assert libro["Talento"]["B4"].value == "Ana García"
    assert libro["Evaluaciones"].max_row == 4  # título, encabezados y 2 evaluaciones (la más reciente arriba)
    assert libro["Evaluaciones"]["C3"].value == 85
    assert libro["Reporte IA"]["D3"].value == EVALUACION["reporte_ia"]


def test_excel_general(general):
    libro = _libro(reportes.excel_general(general, USUARIO, GENERADO))
    assert libro.sheetnames == ["Resumen", "Talentos"]
    # Primera tabla de la hoja (termina en el primer renglón vacío)
    resumen = {}
    for fila in libro["Resumen"].iter_rows(min_row=3, values_only=True):
        if not fila[0]:
            break
        resumen[fila[0]] = fila[1]
    assert resumen["Total de talentos"] == "13" and resumen["Área"] == "Infraestructura/IT/DevOps"
    talentos_hoja = libro["Talentos"]
    assert talentos_hoja.max_row == 15
    encabezados = [celda.value for celda in talentos_hoja[2]]
    assert encabezados[-4:] == ["Puntuación", "Nivel de potencial", "Clasificación 9-box", "Fecha de evaluación"]


def test_textos_de_respaldo(general):
    texto = reportes.texto_individual(_individual(), USUARIO, GENERADO).decode("utf-8-sig")
    assert "Ana García" in texto and "REPORTE DE IA" in texto and reportes.LEYENDA in texto
    texto = reportes.texto_general(general, USUARIO, GENERADO).decode("utf-8-sig")
    assert "INDICADORES" in texto and "Talento 13" in texto


# --- Con MySQL (sin dejar rastro) ---

@pytest.fixture
def evaluado(administrador):
    _, id_talento = talentos.crear(administrador, TALENTO)
    with db.transaccion() as cursor:
        columnas = ("fecha_evaluacion", *evaluaciones.COLUMNAS)
        cursor.execute(
            "INSERT INTO Evaluaciones (id_talento, " + ", ".join(columnas) + ") VALUES (%s, "
            + ", ".join(["%s"] * len(columnas)) + ")",
            (id_talento, datetime.date.today(), *(EVALUACION[c] for c in evaluaciones.COLUMNAS)),
        )
    return id_talento


def _exportaciones(id_usuario):
    with db.transaccion() as cursor:
        cursor.execute("SELECT id_registro_afectado, detalle FROM Bitacora WHERE accion = 'exportacion' "
                       "AND id_usuario = %s ORDER BY id_bitacora DESC LIMIT 1", (id_usuario,))
        fila = cursor.fetchone()
    return fila["id_registro_afectado"], json.loads(fila["detalle"])


def test_exportar_individual_registra_en_bitacora(administrador, evaluado):
    datos = reportes.datos_individual(administrador, evaluado)
    assert datos["ultima"]["puntuacion_talento_por_ia"] == 85 and datos["eventos"][0]["tipo_evento"] == "ingreso"
    resultado = reportes.exportar(administrador, datos, "pdf")
    assert resultado["codigo"] is None and resultado["archivo"].endswith(".pdf")
    assert resultado["contenido"].startswith(b"%PDF")
    id_registro, detalle = _exportaciones(administrador["id_usuario"])
    assert id_registro == evaluado
    assert detalle == {"reporte": "individual", "formato": "pdf", "archivo": resultado["archivo"],
                       "resultado": "generado"}


@pytest.mark.parametrize("formato, codigo", [("pdf", "E-201"), ("excel", "E-202")])
def test_respaldo_en_texto_si_falla(administrador, evaluado, monkeypatch, formato, codigo):
    def falla(*_):
        raise RuntimeError("falla simulada")
    monkeypatch.setitem(reportes._GENERADORES, ("individual", formato), falla)
    resultado = reportes.exportar(administrador, reportes.datos_individual(administrador, evaluado), formato)
    assert resultado["codigo"] == codigo
    assert resultado["archivo"].endswith(".txt") and resultado["mime"] == reportes.MIME_TEXTO
    assert "Ana García" in resultado["contenido"].decode("utf-8-sig")
    assert _exportaciones(administrador["id_usuario"])[1]["resultado"] == codigo


def test_exportar_general_con_filtros(administrador, evaluado):
    datos = reportes.datos_general(administrador, area="Infraestructura/IT/DevOps", niveles=["alto"])
    assert evaluado in {t["id_talento"] for t in datos["talentos"]}
    assert all(t["area"] == "Infraestructura/IT/DevOps" for t in datos["talentos"])
    resultado = reportes.exportar(administrador, datos, "excel")
    assert resultado["codigo"] is None
    id_registro, detalle = _exportaciones(administrador["id_usuario"])
    assert id_registro is None and detalle["filtros"]["niveles"] == ["alto"]


def test_encargado_solo_su_area(administrador, evaluado):
    encargado = {**administrador, "cargo": permisos.ENCARGADO, "area": "Finanzas"}
    with pytest.raises(errores.ErrorSistema) as error:
        reportes.datos_individual(encargado, evaluado)
    assert error.value.codigo == "E-304"
    datos_ajenos = reportes.datos_individual(administrador, evaluado)
    with pytest.raises(errores.ErrorSistema):
        reportes.exportar(encargado, datos_ajenos, "pdf")
    # Aunque pida otra área, el reporte general queda limitado a la suya
    datos = reportes.datos_general(encargado, area="Infraestructura/IT/DevOps")
    assert datos["filtros"]["area"] == "Finanzas"
    assert all(t["area"] == "Finanzas" for t in datos["talentos"])


def test_pantalla_de_reportes(administrador, evaluado, app_con_sesion):
    app = app_con_sesion(administrador).run()
    app.switch_page("paginas/reportes.py").run()
    assert not app.exception
    assert [pestana.label for pestana in app.tabs] == ["Reporte individual", "Reporte general"]
    app.text_input(key="rep_texto").set_value(str(evaluado)).run()
    assert not app.exception and not app.warning
    etiquetas = [boton.label for boton in app.button]
    assert "Generar PDF" in etiquetas and "Generar Excel" in etiquetas  # del reporte general


def test_pantalla_del_encargado(administrador, app_con_sesion):
    encargado = {**administrador, "cargo": permisos.ENCARGADO, "area": "Calidad"}
    app = app_con_sesion(encargado).run()
    app.switch_page("paginas/reportes.py").run()
    assert not app.exception
    fijas = [campo for campo in app.text_input if campo.disabled]
    assert fijas and all(campo.value == "Calidad" for campo in fijas)
