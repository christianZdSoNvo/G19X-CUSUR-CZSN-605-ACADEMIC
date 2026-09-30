import datetime
import random

import pytest
from openpyxl import load_workbook

import datos_ejemplo
from backend import catalogos, db, evaluaciones, ia, talentos

HOY = datetime.date(2026, 9, 30)


@pytest.fixture
def generados():
    return datos_ejemplo.generar(87, random.Random(7), HOY)


def test_datos_validos_para_el_formulario(generados):
    for talento in generados:
        assert talentos.validar(talentos.limpiar(talento)) == {}, talento["nombre"]


def test_nombres_y_correos_unicos(generados):
    assert len({t["nombre"] for t in generados}) == len(generados)
    assert len({t["correo"] for t in generados}) == len(generados)
    assert all(t["correo"].endswith("@correo.com") and t["correo"].isascii() for t in generados)


def test_fechas_y_desempeno_variados(generados):
    assert all(datetime.date(2020, 1, 1) <= t["fecha_ingreso"] <= HOY for t in generados)
    assert len({t["fecha_ingreso"].year for t in generados}) >= 6
    niveles = [ia.nivel_de_puntuacion(t["desempeno_previo"]) for t in generados]
    assert niveles.count("bajo") == niveles.count("medio") == niveles.count("alto") == 29
    assert {t["area"] for t in generados} == set(catalogos.AREAS)


def test_evaluaciones_coherentes(generados):
    todas = [e for t in generados for e in t["evaluaciones"]]
    assert todas
    for talento in generados:
        fechas = [e["fecha_evaluacion"] for e in talento["evaluaciones"]]
        assert fechas == sorted(fechas) and all(talento["fecha_ingreso"] < f <= HOY for f in fechas)
        for evaluacion in talento["evaluaciones"]:
            assert ia.nivel_de_puntuacion(evaluacion["puntuacion_talento_por_ia"]) == \
                evaluacion["potencial_crecimiento_nivel"]
            assert evaluacion["clasificacion_9box"] == evaluaciones.clasificar_9box(
                talento["desempeno_previo"], evaluacion["potencial_crecimiento_nivel"])
            assert len(evaluacion["reporte_ia"]) <= 1000


def test_no_repite_correos_existentes():
    ocupados = {"jose.hernandez@correo.com", "maria.garcia@correo.com"}
    generados = datos_ejemplo.generar(88, random.Random(5), HOY, ocupados)
    assert not ocupados & {t["correo"] for t in generados}


def test_telefonos_en_orden_visual(generados):
    assert [t["tel_mx"] for t in generados[:10]] == [
        "1112345678", "1111234567", "1111123456", "1111112345", "1111111234",
        "1111111123", "1111111112", "1111111111", "2111111111", "2211111111"]
    assert datos_ejemplo.TELEFONOS[17] == "2222222222" and datos_ejemplo.TELEFONOS[18] == "3222222222"
    assert datos_ejemplo.TELEFONOS[-1] == "9999999999" and len(set(datos_ejemplo.TELEFONOS)) == 88
    with pytest.raises(ValueError):
        datos_ejemplo.generar(89, random.Random(1), HOY)


def test_misma_semilla_mismos_datos():
    assert datos_ejemplo.generar(10, random.Random(3), HOY) == datos_ejemplo.generar(10, random.Random(3), HOY)


def test_vista_no_escribe_en_la_bd(administrador, tmp_path):
    with db.transaccion() as cursor:
        cursor.execute("SELECT COUNT(*) AS n FROM Talentos")
        antes = cursor.fetchone()["n"]
    archivo = tmp_path / "vista.xlsx"
    total, _ = datos_ejemplo.vista(8, archivo=archivo)
    libro = load_workbook(archivo)
    assert libro.sheetnames == ["Talentos", "Evaluaciones", "Eventos"]
    assert libro["Talentos"].max_row == total + 1 == 9
    with db.transaccion() as cursor:
        cursor.execute("SELECT COUNT(*) AS n FROM Talentos")
        assert cursor.fetchone()["n"] == antes


def test_cargar_y_borrar(administrador):
    with db.transaccion() as cursor:
        if datos_ejemplo.ids_de_ejemplo(cursor):
            pytest.skip("Ya hay talentos de ejemplo cargados")
        # Un talento real con el mismo dominio no debe borrarse
        cursor.execute("INSERT INTO Talentos (area, nombre, correo, tel_mx, fecha_ingreso, fecha_modificacion) "
                       "VALUES ('Calidad', 'Real', 'real.prueba@correo.com', '5512345678', CURDATE(), CURDATE())")
        id_real = cursor.lastrowid
    total, evaluadas = datos_ejemplo.cargar(12, semilla=1)
    with db.transaccion() as cursor:
        ids = datos_ejemplo.ids_de_ejemplo(cursor)
        assert len(ids) == total == 12
        cursor.execute("SELECT COUNT(*) AS n FROM Evaluaciones WHERE id_talento IN ("
                       + ", ".join(map(str, ids)) + ")")
        assert cursor.fetchone()["n"] == evaluadas
    with pytest.raises(ValueError):
        datos_ejemplo.cargar(5)
    assert datos_ejemplo.borrar() == 12
    with db.transaccion() as cursor:
        assert datos_ejemplo.ids_de_ejemplo(cursor) == []
        cursor.execute("SELECT id_talento FROM Talentos WHERE id_talento = %s", (id_real,))
        assert cursor.fetchone() is not None
        cursor.execute("SELECT COUNT(*) AS n FROM Eventos_talento e LEFT JOIN Talentos t USING (id_talento) "
                       "WHERE t.id_talento IS NULL")
        assert cursor.fetchone()["n"] == 0
